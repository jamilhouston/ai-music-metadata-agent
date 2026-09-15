from __future__ import annotations

import io
import json
import csv

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

from .models import CatalogRecord


def flatten_record(record: CatalogRecord) -> dict[str, str | int | float | None]:
    values: dict[str, str | int | float | None] = {}
    for prefix, model in (("", record.metadata), ("technical_", record.technical), ("rights_", record.rights)):
        for key, value in model.model_dump(mode="json").items():
            if isinstance(value, (list, dict)):
                values[prefix + key] = json.dumps(value, ensure_ascii=False)
            else:
                values[prefix + key] = value
    values.update({
        "source_url": record.source_url,
        "review_status": record.review_status,
        "reviewed_by": record.reviewed_by,
        "reviewed_at": record.reviewed_at.isoformat() if record.reviewed_at else "",
        "notes": record.notes,
    })
    return values


def to_text(record: CatalogRecord) -> str:
    values = flatten_record(record)
    return "\n".join(f"{key.replace('_', ' ').title()}: {value}" for key, value in values.items() if value not in (None, "", "[]", "{}"))


def to_xlsx(record: CatalogRecord) -> bytes:
    values = flatten_record(record)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metadata"
    sheet.append(list(values.keys()))
    sheet.append(list(values.values()))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243447")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def to_csv(record: CatalogRecord) -> bytes:
    """Create a UTF-8 CSV that can be imported directly into Google Sheets."""
    values = flatten_record(record)
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(values.keys())
    writer.writerow("" if value is None else value for value in values.values())
    return output.getvalue().encode("utf-8-sig")


def to_docx(record: CatalogRecord) -> bytes:
    document = Document()
    document.add_heading(record.metadata.track_title or "Music Metadata", 0)
    for key, value in flatten_record(record).items():
        if value not in (None, "", "[]", "{}"):
            document.add_heading(key.replace("_", " ").title(), level=2)
            document.add_paragraph(str(value))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def to_pdf(record: CatalogRecord) -> bytes:
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    story = [Paragraph(record.metadata.track_title or "Music Metadata", styles["Title"]), Spacer(1, 12)]
    rows = []
    for key, value in flatten_record(record).items():
        if value not in (None, "", "[]", "{}"):
            rows.append([Paragraph(key.replace("_", " ").title(), styles["BodyText"]), Paragraph(str(value), styles["BodyText"])])
    table = Table(rows, colWidths=[150, 370], repeatRows=0)
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF0F6")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(table)
    SimpleDocTemplate(output, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36).build(story)
    return output.getvalue()


