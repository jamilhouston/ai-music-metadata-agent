from __future__ import annotations

import re
from io import BytesIO

from openpyxl import load_workbook


BASE_TAXONOMY = {
    "genre_families": [
        "Ambient", "Blues", "Classical", "Country", "Electronic Dance",
        "Reggae", "Folk", "Hip Hop", "Jazz", "Latin", "R&B", "Rock",
        "Pop", "World / Ethnic", "Funk", "Children", "Religious", "Holiday",
    ],
    "subgenres": [
        "Ambient Avant-Garde", "New Age", "Blues Rock", "Classic Blues",
        "Classical Avant-Garde", "Baroque", "Contemporary Classical", "Romantic",
        "Alternative Country", "Bluegrass", "Contemporary Country", "Country Pop",
        "Country Rock", "Traditional Country", "Zydeco", "Big Beat", "Dance Pop",
        "Disco", "Downtempo", "Drum and Bass", "Dubstep", "Future Bass", "Lo-Fi",
        "House", "Indietronica", "New Wave / Synthpop", "Techno", "Trance",
        "Trip Hop", "Classic Reggae", "Dancehall", "Dub", "Classic Folk",
        "Folk Rock", "Indie Folk", "Contemporary Hip Hop", "Southern Rap",
        "East Coast Rap", "Alternative Hip Hop", "Gangsta Rap", "Old School Hip Hop",
        "West Coast Rap / G-Funk", "Trap", "Emo Rap", "Golden Age Hip Hop",
        "Bebop", "Bossa Nova", "Contemporary Jazz", "Cool Jazz", "Dixieland",
        "Jazz Fusion", "Ragtime", "Smooth Jazz", "Swing", "Jazz Pop", "Reggaeton",
        "Tango", "Contemporary R&B", "Doo Wop", "Neo-Soul", "Soul", "Acid Jazz",
        "Alternative R&B", "Alternative Rock", "Garage Rock", "Glam Rock", "Grunge",
        "Hard Rock", "Heartland Rock", "Heavy Metal", "Indie Rock", "Industrial Rock",
        "Pop Rock", "Post-Rock", "Progressive Rock", "Psychedelic Rock",
        "Punk / Pop Punk", "Rockabilly", "Ska Punk", "Soft Rock", "Surf Rock",
    ],
    "mood_families": [
        "Happy", "Sad", "Scary", "Determined", "Energetic", "Relaxed", "Angry",
        "Light", "Spacey", "Suspenseful", "Sexy", "Loving", "Sentimental",
        "Funny", "Elegant", "Weird", "Serious", "Inspirational",
    ],
    "moods": [
        "Bouncy", "Bright", "Celebratory", "Cheerful", "Feel Good", "Joyful",
        "Optimistic", "Playful", "Positive", "Uplifting", "Melancholic", "Somber",
        "Wistful", "Brooding", "Haunting", "Ominous", "Sinister", "Eerie", "Tense",
        "Confident", "Hopeful", "Purposeful", "Driving", "Dynamic", "Exciting",
        "Explosive", "Lively", "Rebellious", "Urgent", "Calm", "Carefree", "Chill",
        "Laid Back", "Meditative", "Mellow", "Peaceful", "Smooth", "Aggressive",
        "Fierce", "Hostile", "Gentle", "Delicate", "Whimsical", "Ethereal",
        "Hypnotic", "Dreamy", "Celestial", "Mysterious", "Investigative", "Foreboding",
        "Sensual", "Sultry", "Seductive", "Romantic", "Tender", "Warm", "Sweet",
    ],
    "style_types": [
        "Performance Based", "Musical Attributes", "Other Description",
        "Film/TV Genres", "Situational Usage", "Playlists",
    ],
    "instrument_families": [
        "Strings", "Winds", "Brass", "Percussion", "Keyboard",
        "Free Reed Aerophone", "Voice", "Electronic",
    ],
    "sync_uses": [
        "Action", "Adventure", "Advertising", "Background", "Comedy", "Crime",
        "Drama", "Educational", "End Credits", "Epic", "Fantasy", "Film Noir",
        "Game Show", "Holiday", "Horror", "Kids Show", "News", "Opening Credits",
        "Promo", "Romance", "Science Fiction", "Sitcom", "Superhero", "Suspense",
        "Thriller", "Trailer", "Underscore", "Western", "Chase", "Corporate",
        "Sports", "Party", "Road Trip", "Vacation",
    ],
}


def clean_term(value: object) -> str:
    if not isinstance(value, str):
        return ""
    value = value.replace("\ufffd", "").replace("\u00a0", " ")
    return re.sub(r"\s+", " ", value).strip()


def unique_terms(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        cleaned = clean_term(value)
        key = cleaned.casefold()
        if cleaned and key not in seen:
            seen.add(key)
            result.append(cleaned)
    return result


def taxonomy_from_workbook(data: bytes) -> dict[str, list[str]]:
    """Read only the vocabulary tabs from a workbook shaped like the user's reference."""
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    taxonomy = {key: list(value) for key, value in BASE_TAXONOMY.items()}
    sheet_names = {name.casefold(): name for name in workbook.sheetnames}

    def values_from(partial_name: str) -> list[list[str]]:
        match = next((name for key, name in sheet_names.items() if partial_name in key), None)
        if not match:
            return []
        return [[clean_term(cell) for cell in row] for row in workbook[match].iter_rows(values_only=True)]

    genres = values_from("genresub genre")
    if genres:
        taxonomy["genre_families"] = unique_terms(genres[0])
        taxonomy["subgenres"] = unique_terms([cell for row in genres[1:] for cell in row])

    moods = values_from("mood")
    if moods:
        taxonomy["mood_families"] = unique_terms(moods[0])
        taxonomy["moods"] = unique_terms([cell for row in moods[1:] for cell in row])

    instruments = values_from("instruments")
    if instruments:
        taxonomy["instrument_families"] = unique_terms(instruments[0])
        taxonomy["instruments"] = unique_terms([cell for row in instruments[1:] for cell in row])

    styles = values_from("style")
    if styles:
        taxonomy["style_types"] = unique_terms(styles[0])
        taxonomy["styles"] = unique_terms([cell for row in styles[1:] for cell in row])

    return taxonomy


def taxonomy_prompt(taxonomy: dict[str, list[str]]) -> str:
    sections = []
    for name, values in taxonomy.items():
        sections.append(f"{name}: {', '.join(values)}")
    return "\n".join(sections)

